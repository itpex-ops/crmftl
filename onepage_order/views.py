from decimal import Decimal, InvalidOperation
from datetime import datetime
import logging
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from .models import (
    Customer,
    Order,
    Tracking,
    TrackingDocument,
    VehiclePayment,
    CustomerPayment,
    TrackingSession,
    LiveLocation,
    SMSLog,
    ApiLog,
    ApiToken,)
from .services.auth_service import TrackingAuthService
from  .services.consent_auth_service import ConsentAuthService
from .services.import_service import ImportService
from .services.consent_service import ConsentService
from .services.location_service import LocationService
from .services.delete_service import DeleteService
from .services.modify_service import ModifyService
from django.utils import timezone
logger = logging.getLogger(__name__)

# =============================================================
# COMMON HELPERS
# =============================================================

def get_post_value(request, field, default=""):
    """Return a cleaned POST string."""
    value = request.POST.get(field, default)
    return value.strip() if isinstance(value, str) else default

def to_decimal(value, default="0.00"):
    """Safely convert POST input to Decimal."""
    if value in (None, ""):
        return Decimal(default)

    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        return Decimal(default)

def to_integer(value, default=0):
    """Safely convert POST input to integer."""
    if value in (None, ""):
        return default

    try:
        return int(value)
    except (ValueError, TypeError):
        return default

def to_date(value):
    """Convert HTML date input (YYYY-MM-DD) to a Python date."""
    if not value:
        return None

    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None

def gst_from_request(request):
    """
    Read and validate the GST checkbox value.

    Supported values:
    0, 7, 8, 10, 13
    """
    raw_value = request.POST.get("gst_percent", "0")
    gst_percent = to_decimal(raw_value, "0.00")

    valid_rates = {
        Decimal("0.00"),
        Decimal("7.00"),
        Decimal("8.00"),
        Decimal("10.00"),
        Decimal("13.00"),
    }

    if gst_percent not in valid_rates:
        raise ValidationError("Invalid GST percentage selected.")

    return gst_percent

def order_queryset():
    """Common optimized Order queryset."""
    return (
        Order.objects
        .select_related(
            "customer",
            "tracking",
            "tracking_session",
        )
        .prefetch_related(
            "vehicle_payments",
            "customer_payments",
        )
        .order_by("-id")
    )

def payment_page_context(orders, payments):
    return {
        "orders": orders,
        "payments": payments,
    }

def add_tracking_template_flags(tracking):
    """
    Backward-compatible template aliases.

    Tracking model stores:
    - live tracking as live_tracking_at
    - balance transfer as balance_to_fleet

    Older tracking_page templates may still use:
    - tracking.live_tracking
    - tracking.balance_trans_fleet
    """
    tracking.live_tracking = bool(
        tracking.live_tracking_at
    )
    tracking.balance_trans_fleet = bool(
        tracking.balance_to_fleet
    )
    return tracking

# =============================================================
# LIVE TRACKING PAGE
# =============================================================
@login_required
def vehicle_live_location(request, pk):
    """
    Return the latest location using a TrackingSession PK.
    """
    tracking_session = get_object_or_404(
        TrackingSession.objects.select_related("order"),
        pk=pk,
    )

    order = tracking_session.order

    latest_location = (
        LiveLocation.objects
        .filter(session=tracking_session)
        .order_by("-received_at")
        .first()
    )

    if not latest_location:
        return JsonResponse({
            "success": True,
            "tracking": True,
            "has_location": False,
            "vehicle_number": order.vehicle_number or "",
            "driver_number": order.driver_number or "",
            "trip_number": order.trip_number or "",
            "tracking_reference": (
                tracking_session.tracking_reference
            ),
            "status": tracking_session.get_status_display(),
            "tracking_enabled": (
                tracking_session.tracking_enabled
            ),
            "consent_received": (
                tracking_session.consent_received
            ),
            "message": "Waiting for vehicle location...",
        })

    latitude = (
        float(latest_location.latitude)
        if latest_location.latitude is not None
        else None
    )

    longitude = (
        float(latest_location.longitude)
        if latest_location.longitude is not None
        else None
    )

    return JsonResponse({
        "success": True,
        "tracking": True,
        "has_location": True,
        "vehicle_number": order.vehicle_number or "",
        "driver_number": order.driver_number or "",
        "trip_number": order.trip_number or "",
        "tracking_reference": (
            tracking_session.tracking_reference
        ),
        "status": tracking_session.get_status_display(),
        "tracking_enabled": (
            tracking_session.tracking_enabled
        ),
        "consent_received": (
            tracking_session.consent_received
        ),
        "latitude": latitude,
        "longitude": longitude,
        "accuracy": latest_location.accuracy,
        "location_name": (
            latest_location.location_name or ""
        ),
        "address": latest_location.address or "",
        "location_status": (
            latest_location.location_status or ""
        ),
        "tracked": latest_location.tracked,
        "received_at": (
            latest_location.received_at.isoformat()
        ),
        "session_last_updated": (
            tracking_session.last_updated.isoformat()
            if tracking_session.last_updated
            else None
        ),
    })

# =============================================================
# ORDER DELETE
# =============================================================

@login_required
def onepageorder_delete(request, pk):
    if request.method != "POST":
        messages.error(
            request,
            "Invalid request.",
        )
        return redirect("onepageorder_list")

    order = get_object_or_404(
        Order,
        pk=pk,
    )

    trip_number = order.trip_number

    try:
        with transaction.atomic():
            order.delete()

        messages.success(
            request,
            f"Order {trip_number} deleted successfully.",
        )

    except Exception:
        logger.exception(
            "Unable to delete order %s",
            trip_number,
        )
        messages.error(
            request,
            f"Unable to delete order {trip_number}.",
        )

    return redirect("onepageorder_list")

def excel_datetime(value):
    """
    Convert Django timezone-aware datetime to
    timezone-naive datetime for Excel.
    """

    if value is None:
        return None

    if hasattr(value, "tzinfo"):

        if value.tzinfo is not None:

            value = timezone.localtime(value)

            value = value.replace(
                tzinfo=None
            )

    return value

# Backward-compatible name if an old URL still points to delete_vehicle.
@login_required
def delete_vehicle(request, pk):
    return onepageorder_delete(request, pk)

# =============================================================
# ORDER LIST
# =============================================================

@login_required
def onepageorder_list(request):
    search = request.GET.get("q", "").strip()

    orders = order_queryset()

    if search:
        orders = orders.filter(
            Q(trip_number__icontains=search)
            | Q(customer__name__icontains=search)
            | Q(origin__icontains=search)
            | Q(destination__icontains=search)
            | Q(vehicle_number__icontains=search)
            | Q(vehicle_type__icontains=search)
        )

    return render(
        request,
        "onepageorders/order_list.html",
        {
            "orders": orders,
            "q": search,
        },
    )

# =============================================================
# ORDER CREATE HELPERS
# =============================================================

def create_customer_from_request(request):
    customer_name = get_post_value(
        request,
        "customer_name",
    )

    if not customer_name:
        raise ValidationError(
            "Customer Name is required."
        )

    customer = Customer(
        name=customer_name,
        contact_number=get_post_value(
            request,
            "contact_number",
        ),
        email=get_post_value(
            request,
            "email",
        ),
        address=get_post_value(
            request,
            "address",
        ),
    )

    customer.full_clean()
    customer.save()

    return customer


def build_order_from_request(request, customer):
    """
    Build an Order from POST data.
    The object is returned without saving.
    """

    order = Order(
        customer=customer,

        # Customer / Sales
        lead_generated_through=get_post_value(
            request,
            "lead_generated_through",
        ),
        reference_name=get_post_value(
            request,
            "reference_name",
        ),
        sales_closed_by=get_post_value(
            request,
            "sales_closed_by",
        ),

        # Shipment
        origin=get_post_value(
            request,
            "origin",
        ),
        destination=get_post_value(
            request,
            "destination",
        ),
        material=get_post_value(
            request,
            "material",
        ),
        packing_type=get_post_value(
            request,
            "packing_type",
        ),
        no_of_pieces=to_integer(
            request.POST.get("no_of_pieces"),
        ),
        weight_tons=to_decimal(
            request.POST.get("weight_tons"),
            "0.000",
        ),

        # Vehicle
        vehicle_type=get_post_value(
            request,
            "vehicle_type",
        ),
        vehicle_number=get_post_value(
            request,
            "vehicle_number",
        ),
        driver_number=get_post_value(
            request,
            "driver_number",
        ),
        owner_number=get_post_value(
            request,
            "owner_number",
        ),
        vehicle_sourced_by=request.POST.get(
            "vehicle_sourced_by",
            "Direct",
        ),
        owner_broker_name=get_post_value(
            request,
            "owner_broker_name",
        ),

        # Commercials
        freight_amount=to_decimal(
            request.POST.get("freight_amount"),
        ),
        loading_unloading_charges=to_decimal(
            request.POST.get(
                "loading_unloading_charges"
            ),
        ),
        halting_charges=to_decimal(
            request.POST.get(
                "halting_charges"
            ),
        ),
        other_charges=to_decimal(
            request.POST.get("other_charges"),
        ),
        selling_amount=to_decimal(
            request.POST.get("selling_amount"),
        ),
        gst_percent=gst_from_request(request),

        manager_approval=request.POST.get(
            "manager_approval",
            "Pending",
        ),
        approved_by=get_post_value(
            request,
            "approved_by",
        ),

        # Customer Payment Terms
        customer_billing_type=request.POST.get(
            "customer_billing_type",
            "",
        ),
        customer_payment_type=request.POST.get(
            "customer_payment_type",
            "",
        ),
        customer_advance_amount=to_decimal(
            request.POST.get(
                "customer_advance_amount"
            ),
        ),
        customer_balance_amount=to_decimal(
            request.POST.get(
                "customer_balance_amount"
            ),
        ),
        customer_payment_method=request.POST.get(
            "customer_payment_method",
            "",
        ),
        promised_due_date=to_date(
            request.POST.get(
                "promised_due_date"
            )
        ),

        # Contracted Vehicle Payment Terms
        vehicle_advance_amount=to_decimal(
            request.POST.get(
                "vehicle_advance_amount"
            ),
        ),
        vehicle_balance_amount=to_decimal(
            request.POST.get(
                "vehicle_balance_amount"
            ),
        ),
        vehicle_owner_name=get_post_value(
            request,
            "vehicle_owner_name",
        ),
        pan_card=get_post_value(
            request,
            "pan_card",
        ),
        account_name=get_post_value(
            request,
            "account_name",
        ),
        account_number=get_post_value(
            request,
            "account_number",
        ),
        ifsc_code=get_post_value(
            request,
            "ifsc_code",
        ),
        upi_number=get_post_value(
            request,
            "upi_number",
        ),

        # Options
        send_sms=(
            request.POST.get("send_sms") == "on"
        ),
        create_agreement_tds=(
            request.POST.get(
                "create_agreement_tds"
            ) == "on"
        ),
    )

    return order


def prepare_order_for_validation(order):
    """
    Order.full_clean() runs before Order.save(), so generate
    trip_number first.
    """
    if not order.trip_number:
        order.trip_number = Order.generate_trip_number()

    return order

# =============================================================
# ORDER CREATE
# =============================================================

@login_required
def onepageorder_create(request):
    template_name = "onepageorders/order_create.html"
    if request.method != "POST":
        return render(
            request,
            template_name,
        )
    try:
        with transaction.atomic():

            customer = create_customer_from_request(
                request
            )

            order = build_order_from_request(
                request,
                customer,
            )

            prepare_order_for_validation(
                order
            )

            order.full_clean()
            order.save()

            # Create the workflow record together with the order.
            Tracking.objects.get_or_create(
                order=order
            )

        messages.success(
            request,
            (
                f"Order {order.trip_number} "
                "created successfully."
            ),
        )

        return redirect(
            "onepageorder_detail",
            pk=order.pk,
        )

    except ValidationError as e:
        logger.warning(
            "Order validation failed: %s",
            e,
            exc_info=True,
        )

        if hasattr(e, "message_dict"):
            error_messages = []

            for field, errors in (
                e.message_dict.items()
            ):
                for error in errors:
                    error_messages.append(
                        f"{field}: {error}"
                    )

            messages.error(
                request,
                " | ".join(error_messages),
            )
        else:
            messages.error(
                request,
                str(e),
            )

        return render(
            request,
            template_name,
        )

    except Exception:
        logger.exception(
            "Unexpected error while creating order"
        )

        messages.error(
            request,
            (
                "Unable to create the order. "
                "Please check the entered details "
                "and try again."
            ),
        )

        return render(
            request,
            template_name,
        )

# =============================================================
# ORDER DETAIL
# =============================================================

@login_required
def onepageorder_detail(request, pk):
    order = get_object_or_404(
        Order.objects
        .select_related(
            "customer",
            "tracking",
            "tracking_session",
        )
        .prefetch_related(
            "vehicle_payments",
            "customer_payments",
            "tracking__documents",
        ),
        pk=pk,
    )

    return render(
        request,
        "onepageorders/order_detail.html",
        {
            "order": order,
            "tracking": getattr(
                order,
                "tracking",
                None,
            ),
        },
    )

# =============================================================
# LIVE TRACKING - ORDER LOCATION API
# =============================================================

@login_required
def order_live_location(request, pk):
    order = get_object_or_404(
        Order.objects.select_related(
            "tracking_session",
        ),
        pk=pk,
    )

    tracking_session = getattr(
        order,
        "tracking_session",
        None,
    )

    if not tracking_session:
        return JsonResponse(
            {
                "success": False,
                "tracking": False,
                "message": (
                    "Live tracking is not enabled "
                    "for this order."
                ),
            },
            status=404,
        )

    latest_location = (
        LiveLocation.objects
        .filter(session=tracking_session)
        .order_by("-received_at")
        .first()
    )

    if not latest_location:
        return JsonResponse({
            "success": True,
            "tracking": True,
            "has_location": False,
            "trip_number": order.trip_number or "",
            "vehicle_number": (
                order.vehicle_number or ""
            ),
            "driver_number": (
                order.driver_number or ""
            ),
            "status": (
                tracking_session.get_status_display()
            ),
            "tracking_reference": (
                tracking_session.tracking_reference
            ),
            "driver_mobile": (
                tracking_session.driver_mobile
            ),
            "tracking_enabled": (
                tracking_session.tracking_enabled
            ),
            "consent_received": (
                tracking_session.consent_received
            ),
            "message": (
                "Waiting for vehicle location..."
            ),
        })

    latitude = (
        float(latest_location.latitude)
        if latest_location.latitude is not None
        else None
    )

    longitude = (
        float(latest_location.longitude)
        if latest_location.longitude is not None
        else None
    )

    return JsonResponse({
        "success": True,
        "tracking": True,
        "has_location": True,
        "trip_number": order.trip_number or "",
        "vehicle_number": (
            order.vehicle_number or ""
        ),
        "driver_number": (
            order.driver_number or ""
        ),
        "status": (
            tracking_session.get_status_display()
        ),
        "tracking_enabled": (
            tracking_session.tracking_enabled
        ),
        "consent_received": (
            tracking_session.consent_received
        ),
        "tracking_reference": (
            tracking_session.tracking_reference
        ),
        "driver_mobile": (
            tracking_session.driver_mobile
        ),
        "latitude": latitude,
        "longitude": longitude,
        "accuracy": latest_location.accuracy,
        "location_name": (
            latest_location.location_name or ""
        ),
        "address": (
            latest_location.address or ""
        ),
        "location_status": (
            latest_location.location_status or ""
        ),
        "tracked": latest_location.tracked,
        "received_at": (
            latest_location.received_at.isoformat()
        ),
        "session_last_updated": (
            tracking_session.last_updated.isoformat()
            if tracking_session.last_updated
            else None
        ),
    })


# =============================================================
# TRACKING WORKFLOW
# =============================================================

@login_required
def tracking_page(request, pk):
    """
    Order Tracking Page.

    Tracking workflow fields:
        vehicle_placed
        vehicle_document
        invoice_eway
        advance_to_fleet
        balance_to_fleet
        fleet_departed
        arrived
        delivered
        pod_received
        lr_no_b
        settled

    Live Tracking:
        Tracking.live_tracking_at = timestamp/event
        TrackingSession = actual SmartTrail tracking session
    """

    order = get_object_or_404(
        Order.objects.select_related(
            "customer",
            "tracking",
            "tracking_session",
        ),
        pk=pk,
    )

    # ------------------------------------------------------------
    # GET / CREATE TRACKING
    # ------------------------------------------------------------

    tracking, _created = Tracking.objects.get_or_create(
        order=order
    )

    add_tracking_template_flags(tracking)

    # ------------------------------------------------------------
    # POST
    # ------------------------------------------------------------

    if request.method == "POST":

        # --------------------------------------------------------
        # LOCK AFTER SETTLEMENT
        # --------------------------------------------------------

        if tracking.settled:
            messages.warning(
                request,
                "Tracking already settled. Editing is locked.",
            )

            return redirect(
                "onepageorder_detail",
                pk=order.pk,
            )

        try:
            with transaction.atomic():

                now = timezone.now()

                # =================================================
                # CHECKBOXES
                # =================================================

                tracking.vehicle_document = (
                    "vehicle_document" in request.POST
                )

                tracking.vehicle_placed = (
                    "vehicle_placed" in request.POST
                )

                tracking.invoice_eway = (
                    "invoice_eway" in request.POST
                )

                tracking.advance_to_fleet = (
                    "advance_to_fleet" in request.POST
                )

                # HTML field:
                # balance_trans_fleet
                #
                # Model field:
                # balance_to_fleet

                tracking.balance_to_fleet = (
                    "balance_trans_fleet" in request.POST
                )

                tracking.fleet_departed = (
                    "fleet_departed" in request.POST
                )

                tracking.arrived = (
                    "arrived" in request.POST
                )

                tracking.delivered = (
                    "delivered" in request.POST
                )

                tracking.pod_received = (
                    "pod_received" in request.POST
                )

                tracking.settled = (
                    "settled" in request.POST
                )

                # =================================================
                # LR
                # =================================================

                tracking.lr_no_b = (
                    "lr_no_b" in request.POST
                )

                tracking.lr_no = get_post_value(
                    request,
                    "lr_no",
                )

                # =================================================
                # REMARKS
                # =================================================

                tracking.remarks = get_post_value(
                    request,
                    "remarks",
                )

                # =================================================
                # LIVE TRACKING
                # =================================================
                #
                # live_tracking_at is an event timestamp.
                #
                # Once enabled, DO NOT remove it when later
                # workflow statuses are completed.
                # =================================================

                live_tracking_requested = (
                    "live_tracking" in request.POST
                )

                if (
                    live_tracking_requested
                    and not tracking.live_tracking_at
                ):
                    tracking.live_tracking_at = now

                # =================================================
                # TIMELINE
                # =================================================

                if (
                    tracking.vehicle_placed
                    and not tracking.vehicle_placed_at
                ):
                    tracking.vehicle_placed_at = now

                if (
                    tracking.fleet_departed
                    and not tracking.fleet_departed_at
                ):
                    tracking.fleet_departed_at = now

                if (
                    tracking.arrived
                    and not tracking.arrived_at
                ):
                    tracking.arrived_at = now

                if (
                    tracking.delivered
                    and not tracking.delivered_at
                ):
                    tracking.delivered_at = now

                # =================================================
                # CURRENT WORKFLOW STATUS
                # =================================================
                #
                # Live Tracking is NOT used as the permanent
                # current status once the trip progresses.
                #
                # Example:
                # Live Tracking enabled
                #        ↓
                # Fleet Departed
                #        ↓
                # Arrived
                #        ↓
                # Delivered
                #
                # live_tracking_at remains stored.
                # =================================================

                if tracking.settled:
                    tracking.status = "settled"

                elif tracking.pod_received:
                    tracking.status = "pod_received"

                elif tracking.delivered:
                    tracking.status = "delivered"

                elif tracking.arrived:
                    tracking.status = "arrived"

                elif tracking.balance_to_fleet:
                    tracking.status = "balance_trans_fleet"

                elif tracking.fleet_departed:
                    tracking.status = "fleet_departed"

                elif tracking.advance_to_fleet:
                    tracking.status = "advance_to_fleet"

                elif tracking.invoice_eway:
                    tracking.status = "invoice_eway"

                elif tracking.lr_no_b:
                    tracking.status = "lr_generated"

                elif tracking.vehicle_document:
                    tracking.status = "vehicle_document"

                elif tracking.live_tracking_at:
                    tracking.status = "live_tracking"

                elif tracking.vehicle_placed:
                    tracking.status = "vehicle_placed"

                tracking.save()

                add_tracking_template_flags(tracking)

                # =================================================
                # DOCUMENT UPLOAD
                # =================================================

                uploaded_files = request.FILES.getlist(
                    "documents"
                )

                for uploaded_file in uploaded_files:

                    TrackingDocument.objects.create(
                        tracking=tracking,
                        file=uploaded_file,
                    )

            # ====================================================
            # LIVE TRACKING NAVIGATION
            # ====================================================

            if live_tracking_requested:

                tracking_session = getattr(
                    order,
                    "tracking_session",
                    None,
                )

                if tracking_session:

                    return redirect(
                        "onepageordervehicle_live",
                        pk=tracking_session.pk,
                    )

                return redirect(
                    "onepageorder_import_driver",
                    pk=order.pk,
                )
            # ====================================================
            # NORMAL SAVE
            # ====================================================

            messages.success(
                request,
                "Tracking updated successfully.",
            )

            return redirect(
                "onepageorder_detail",
                pk=order.pk,
            )

        except ValidationError as e:

            logger.warning(
                "Tracking validation failed: %s",
                e,
                exc_info=True,
            )

            messages.error(
                request,
                str(e),
            )

        except Exception:

            logger.exception(
                "Unexpected error while updating tracking "
                "for order %s",
                order.pk,
            )

            messages.error(
                request,
                "Unable to update tracking. Please try again.",
            )

    # ------------------------------------------------------------
    # LIVE TRACKING SESSION
    # ------------------------------------------------------------

    tracking_session = getattr(
        order,
        "tracking_session",
        None,
    )

    # ------------------------------------------------------------
    # RENDER
    # ------------------------------------------------------------

    return render(
        request,
        "onepageorders/tracking.html",
        {
            "order": order,
            "tracking": tracking,
            "tracking_session": tracking_session,
            "documents": tracking.documents.all(),
            "live_tracking_enabled": bool(
                tracking.live_tracking_at
            ),
        },
    )

# =============================================================
# ORDER EDIT
# =============================================================

@login_required
def onepageorder_edit(request, pk):

    order = get_object_or_404(
        Order.objects.select_related("customer"),
        pk=pk,
    )

    if request.method == "POST":

        try:
            with transaction.atomic():

                # ------------------------------------------------
                # CUSTOMER
                # ------------------------------------------------

                customer = order.customer

                customer.name = get_post_value(
                    request,
                    "customer_name",
                )

                customer.contact_number = get_post_value(
                    request,
                    "contact_number",
                )

                customer.email = get_post_value(
                    request,
                    "email",
                )

                customer.address = get_post_value(
                    request,
                    "address",
                )

                if not customer.name:
                    raise ValidationError(
                        "Customer Name is required."
                    )

                customer.full_clean()
                customer.save()

                # ------------------------------------------------
                # CUSTOMER / SALES
                # ------------------------------------------------

                order.lead_generated_through = get_post_value(
                    request,
                    "lead_generated_through",
                )

                order.reference_name = get_post_value(
                    request,
                    "reference_name",
                )

                order.sales_closed_by = get_post_value(
                    request,
                    "sales_closed_by",
                )

                # ------------------------------------------------
                # SHIPMENT
                # ------------------------------------------------

                order.origin = get_post_value(
                    request,
                    "origin",
                )

                order.destination = get_post_value(
                    request,
                    "destination",
                )

                order.material = get_post_value(
                    request,
                    "material",
                )

                order.packing_type = get_post_value(
                    request,
                    "packing_type",
                )

                order.no_of_pieces = to_integer(
                    request.POST.get(
                        "no_of_pieces"
                    )
                )

                order.weight_tons = to_decimal(
                    request.POST.get(
                        "weight_tons"
                    ),
                    "0.000",
                )

                # ------------------------------------------------
                # VEHICLE
                # ------------------------------------------------

                order.vehicle_type = get_post_value(
                    request,
                    "vehicle_type",
                )

                order.vehicle_number = get_post_value(
                    request,
                    "vehicle_number",
                )

                order.driver_number = get_post_value(
                    request,
                    "driver_number",
                )

                order.owner_number = get_post_value(
                    request,
                    "owner_number",
                )

                order.vehicle_sourced_by = request.POST.get(
                    "vehicle_sourced_by",
                    "Direct",
                )

                order.owner_broker_name = get_post_value(
                    request,
                    "owner_broker_name",
                )

                # ------------------------------------------------
                # COMMERCIALS
                # ------------------------------------------------

                order.freight_amount = to_decimal(
                    request.POST.get(
                        "freight_amount"
                    )
                )

                order.loading_unloading_charges = to_decimal(
                    request.POST.get(
                        "loading_unloading_charges"
                    )
                )

                order.halting_charges = to_decimal(
                    request.POST.get(
                        "halting_charges"
                    )
                )

                order.other_charges = to_decimal(
                    request.POST.get(
                        "other_charges"
                    )
                )

                order.selling_amount = to_decimal(
                    request.POST.get(
                        "selling_amount"
                    )
                )

                order.gst_percent = gst_from_request(
                    request
                )

                order.manager_approval = request.POST.get(
                    "manager_approval",
                    "Pending",
                )

                order.approved_by = get_post_value(
                    request,
                    "approved_by",
                )

                # ------------------------------------------------
                # CUSTOMER PAYMENT TERMS
                # ------------------------------------------------

                order.customer_billing_type = request.POST.get(
                    "customer_billing_type",
                    "",
                )

                order.customer_payment_type = request.POST.get(
                    "customer_payment_type",
                    "",
                )

                order.customer_advance_amount = to_decimal(
                    request.POST.get(
                        "customer_advance_amount"
                    )
                )

                order.customer_balance_amount = to_decimal(
                    request.POST.get(
                        "customer_balance_amount"
                    )
                )

                order.customer_payment_method = request.POST.get(
                    "customer_payment_method",
                    "",
                )

                order.promised_due_date = to_date(
                    request.POST.get(
                        "promised_due_date"
                    )
                )

                # ------------------------------------------------
                # CONTRACTED VEHICLE PAYMENT
                # ------------------------------------------------

                order.vehicle_advance_amount = to_decimal(
                    request.POST.get(
                        "vehicle_advance_amount"
                    )
                )

                order.vehicle_balance_amount = to_decimal(
                    request.POST.get(
                        "vehicle_balance_amount"
                    )
                )

                order.vehicle_owner_name = get_post_value(
                    request,
                    "vehicle_owner_name",
                )

                order.pan_card = get_post_value(
                    request,
                    "pan_card",
                )

                order.account_name = get_post_value(
                    request,
                    "account_name",
                )

                order.account_number = get_post_value(
                    request,
                    "account_number",
                )

                order.ifsc_code = get_post_value(
                    request,
                    "ifsc_code",
                )

                order.upi_number = get_post_value(
                    request,
                    "upi_number",
                )

                # ------------------------------------------------
                # OPTIONS
                # ------------------------------------------------

                order.send_sms = (
                    request.POST.get("send_sms")
                    == "on"
                )

                order.create_agreement_tds = (
                    request.POST.get(
                        "create_agreement_tds"
                    )
                    == "on"
                )

                # ------------------------------------------------
                # VALIDATE + SAVE
                # ------------------------------------------------

                order.full_clean()
                order.save()

            messages.success(
                request,
                (
                    f"Order {order.trip_number} "
                    "updated successfully."
                ),
            )

            return redirect(
                "onepageorder_detail",
                pk=order.pk,
            )

        except ValidationError as e:

            logger.warning(
                "Order edit validation failed: %s",
                e,
                exc_info=True,
            )

            if hasattr(e, "message_dict"):
                error_messages = []

                for field, errors in (
                    e.message_dict.items()
                ):
                    for error in errors:
                        error_messages.append(
                            f"{field}: {error}"
                        )

                messages.error(
                    request,
                    " | ".join(error_messages),
                )
            else:
                messages.error(
                    request,
                    str(e),
                )

        except Exception:

            logger.exception(
                "Unexpected error while editing order %s",
                order.pk,
            )

            messages.error(
                request,
                "Unable to update order. Please try again.",
            )

    return render(
        request,
        "onepageorders/order_edit.html",
        {
            "order": order,
        },
    )


# =============================================================
# VEHICLE PAYMENTS
# =============================================================

@login_required
def vehicle_payments(request):

    search = request.GET.get("q", "").strip()

    orders = order_queryset()

    if search:
        orders = orders.filter(
            Q(trip_number__icontains=search)
            | Q(customer__name__icontains=search)
            | Q(origin__icontains=search)
            | Q(destination__icontains=search)
            | Q(vehicle_number__icontains=search)
            | Q(vehicle_type__icontains=search)
        )

    payments = (
        VehiclePayment.objects
        .select_related(
            "order",
            "order__customer",
        )
        .order_by(
            "-paid_at",
            "-id",
        )
    )

    if request.method == "POST":

        try:
            pk = get_post_value(
                request,
                "order",
            )

            vehicle_number = get_post_value(
                request,
                "vehicle_number",
            ).upper()

            payment_type = get_post_value(
                request,
                "payment_type",
            )

            amount_value = get_post_value(
                request,
                "amount",
            )

            transaction_reference = get_post_value(
                request,
                "transaction_reference",
            )

            if not pk:
                raise ValidationError(
                    "Please select a Trip / Order."
                )

            order = get_object_or_404(
                Order,
                pk=pk,
            )

            if not vehicle_number:
                raise ValidationError(
                    "Vehicle Number is required."
                )

            valid_payment_types = {
                "Advance",
                "Balance",
                "Others",
            }

            if payment_type not in valid_payment_types:
                raise ValidationError(
                    "Please select a valid Payment Type."
                )

            amount = to_decimal(
                amount_value,
                "0.00",
            )

            if amount <= 0:
                raise ValidationError(
                    "Payment amount must be greater than zero."
                )

            with transaction.atomic():
                payment = VehiclePayment.objects.create(
                    order=order,
                    vehicle_number=vehicle_number,
                    payment_type=payment_type,
                    amount=amount,
                    transaction_reference=(
                        transaction_reference
                    ),
                )

            messages.success(
                request,
                (
                    f"Vehicle payment of "
                    f"₹{payment.amount:,.2f} "
                    "saved successfully."
                ),
            )

            return redirect(
                "vehicle_payments"
            )

        except ValidationError as e:

            messages.error(
                request,
                str(e),
            )

        except Exception:

            logger.exception(
                "Unexpected vehicle payment error"
            )

            messages.error(
                request,
                "Unable to save vehicle payment.",
            )

    return render(
        request,
        "onepageorders/vehicle_payments.html",
        payment_page_context(
            orders,
            payments,
        ),
    )


# =============================================================
# CUSTOMER PAYMENTS
# =============================================================

@login_required
def customer_payments(request):

    search = request.GET.get("q", "").strip()

    orders = order_queryset()

    if search:
        orders = orders.filter(
            Q(trip_number__icontains=search)
            | Q(customer__name__icontains=search)
            | Q(origin__icontains=search)
            | Q(destination__icontains=search)
            | Q(vehicle_number__icontains=search)
            | Q(vehicle_type__icontains=search)
        )

    payments = (
        CustomerPayment.objects
        .select_related(
            "order",
            "order__customer",
        )
        .order_by(
            "-received_at",
            "-id",
        )
    )

    if request.method == "POST":

        try:
            pk = get_post_value(
                request,
                "order",
            )

            received_amount = to_decimal(
                get_post_value(
                    request,
                    "received_amount",
                ),
                "0.00",
            )

            received_through = get_post_value(
                request,
                "received_through",
            )

            utr_details = get_post_value(
                request,
                "utr_details",
            )

            if not pk:
                raise ValidationError(
                    "Please select a Trip / Order."
                )

            order = get_object_or_404(
                Order,
                pk=pk,
            )

            # Customer payment is against the GST-inclusive
            # total selling amount.
            selling_amount = order.total_selling_amount

            if selling_amount <= 0:
                raise ValidationError(
                    (
                        "Total selling amount is not available "
                        f"for Trip {order.trip_number}."
                    )
                )

            if received_amount < 0:
                raise ValidationError(
                    "Received Amount cannot be negative."
                )

            valid_modes = {
                "RTGS",
                "NEFT",
                "CASH",
                "IMPS",
                "UPI",
            }

            if received_through not in valid_modes:
                raise ValidationError(
                    "Please select a valid payment mode."
                )

            with transaction.atomic():
                payment = CustomerPayment.objects.create(
                    order=order,
                    selling_amount=selling_amount,
                    received_amount=received_amount,
                    received_through=received_through,
                    utr_details=utr_details,
                )

            messages.success(
                request,
                (
                    f"Customer payment of "
                    f"₹{payment.received_amount:,.2f} "
                    "saved successfully."
                ),
            )

            return redirect(
                "customer_payments"
            )

        except ValidationError as e:

            messages.error(
                request,
                str(e),
            )

        except Exception:

            logger.exception(
                "Unexpected customer payment error"
            )

            messages.error(
                request,
                "Unable to save customer payment.",
            )

    return render(
        request,
        "onepageorders/customer_payments.html",
        payment_page_context(
            orders,
            payments,
        ),
    )


# =============================================================
# ADMIN MARGIN
# =============================================================

@login_required
def admin_margin(request):

    orders = (
        Order.objects
        .select_related(
            "customer",
            "tracking",
        )
        .prefetch_related(
            "vehicle_payments",
            "customer_payments",
        )
        .order_by("-id")
    )

    return render(
        request,
        "onepageorders/admin_margin.html",
        {
            "orders": orders,
        },
    )

# =============================================================
# HELPER FUNCTIONS
# =============================================================

def get_driver_mobile(session):
    """
    Get driver mobile from TrackingSession.

    Falls back to Order.driver_number when session mobile
    is not available.
    """
    return (
        session.driver_mobile
        or session.order.driver_number
        or ""
    )


def unique_location_history(session):
    """
    Return unique location history.

    Latest record for each latitude/longitude is retained.
    """
    locations = (
        session.locations
        .all()
        .order_by("-received_at")
    )

    history = []
    seen = set()

    for location in locations:

        key = (
            location.latitude,
            location.longitude,
        )

        if key in seen:
            continue

        seen.add(key)
        history.append(location)

    return history


# =============================================================
# LIVE TRACKING LIST
# =============================================================

from django.contrib.auth.decorators import login_required
from django.db.models import Q, Prefetch
from django.shortcuts import render

from .models import Order, LiveLocation


from django.contrib.auth.decorators import login_required
from django.db.models import Q, Prefetch
from django.shortcuts import render

from .models import Order, LiveLocation


@login_required
def onepageorder_tracking_list(request):

    query = request.GET.get("q", "").strip()

    # =========================================================
    # LOCATION HISTORY
    # Latest location first
    # =========================================================

    location_history = LiveLocation.objects.order_by(
        "-received_at"
    )

    # =========================================================
    # ORDERS
    # =========================================================

    orders = (
        Order.objects
        .select_related(
            "customer",
            "tracking",
            "tracking_session",
        )
        .prefetch_related(
            Prefetch(
                "tracking_session__locations",
                queryset=location_history,
                to_attr="tracking_history",
            )
        )
        .order_by("-id")
    )

    # =========================================================
    # SEARCH
    # =========================================================

    if query:

        orders = orders.filter(
            Q(trip_number__icontains=query)
            | Q(customer__name__icontains=query)
            | Q(vehicle_number__icontains=query)
            | Q(driver_number__icontains=query)
            | Q(origin__icontains=query)
            | Q(destination__icontains=query)
        )

    # =========================================================
    # RESPONSE
    # =========================================================

    return render(
        request,
        "onepageorders/list.html",
        {
            "order": orders,
            "q": query,
        },
    )
# =============================================================
# LIVE TRACKING SETUP
# pk = ORDER PK
# =============================================================

@login_required
def onepageorder_tracking_setup(request, pk):

    order = get_object_or_404(
        Order.objects.select_related(
            "customer",
            "tracking",
            "tracking_session",
        ),
        pk=pk,
    )

    session = getattr(
        order,
        "tracking_session",
        None,
    )

    return render(
        request,
        "onepageorders/setup.html",
        {
            "order": order,
            "session": session,
        },
    )


# =============================================================
# IMPORT DRIVER
# pk = ORDER PK
# =============================================================

@login_required
def import_driver(request, pk):

    order = get_object_or_404(
        Order,
        pk=pk,
    )

    result = ImportService.import_driver(
        order
    )

    if result.get("success"):

        messages.success(
            request,
            (
                "Driver imported successfully "
                "into Telenity."
            ),
        )

        return redirect(
            "onepageorder_tracking_setup",
            pk=order.pk,
        )

    # ---------------------------------------------------------
    # ERROR MESSAGE
    # ---------------------------------------------------------

    error_message = result.get(
        "message",
        "Driver import failed.",
    )

    if isinstance(
        error_message,
        dict,
    ):

        error_message = (
            error_message.get("errorMessage")
            or error_message.get("message")
            or error_message.get("raw_response")
            or str(error_message)
        )

    messages.error(
        request,
        str(error_message),
    )

    return redirect(
        "onepageorder_tracking_setup",
        pk=order.pk,
    )


# =============================================================
# SEND CONSENT
# pk = TRACKING SESSION PK
# =============================================================

# =============================================================
# SEND CONSENT
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def onepageorder_send_consent(request, pk):

    if request.method != "POST":

        messages.warning(
            request,
            "Invalid request method.",
        )

        return redirect(
            "onepageorder_tracking_setup",
            pk=pk,
        )

    session = get_object_or_404(
        TrackingSession,
        pk=pk,
    )

    result = ConsentService.onepageorder_send_consent(
        session
    )

    if result.get("success"):

        messages.success(
            request,
            "Consent SMS sent successfully.",
        )

    else:

        message = result.get(
            "message",
            "Unable to send consent SMS.",
        )

        if isinstance(message, dict):

            message = (
                message.get("errorMessage")
                or message.get("message")
                or message.get("raw_response")
                or str(message)
            )

        messages.error(
            request,
            str(message),
        )

    return redirect(
        "onepageorder_tracking_setup",
        pk=session.order.pk,
    )


# =============================================================
# CHECK CONSENT
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def onepageorder_check_consent(request, pk):

    if request.method not in (
        "GET",
        "POST",
    ):

        messages.warning(
            request,
            "Invalid request method.",
        )

        return redirect(
            "onepageorder_tracking_list"
        )

    session = get_object_or_404(
        TrackingSession,
        pk=pk,
    )

    result = ConsentService.onepageorder_check_consent(
        session
    )

    if not result.get("success"):

        message = result.get(
            "message",
            "Unable to check consent.",
        )

        if isinstance(message, dict):

            message = (
                message.get("errorMessage")
                or message.get("message")
                or message.get("raw_response")
                or str(message)
            )

        messages.error(
            request,
            str(message),
        )

        return redirect(
            "onepageorder_tracking_setup",
            pk=session.order.pk,
        )

    # Service has already updated the session.
    session.refresh_from_db()

    if (
        session.consent_received
        and session.tracking_enabled
    ):

        messages.success(
            request,
            (
                "Consent received. "
                "Live tracking has been activated."
            ),
        )

    elif session.status == "sms_sent":

        messages.warning(
            request,
            "Driver consent is still pending.",
        )

    elif session.status == "license_hold":

        messages.warning(
            request,
            "Tracking is currently on license hold.",
        )

    elif session.status == "expired":

        messages.warning(
            request,
            "Tracking consent has expired.",
        )

    elif session.status == "error":

        messages.error(
            request,
            (
                "Consent was received, but "
                "tracking could not be activated."
            ),
        )

    else:

        messages.info(
            request,
            (
                "Tracking status: "
                f"{session.get_status_display()}"
            ),
        )

    return redirect(
        "onepageorder_tracking_setup",
        pk=session.order.pk,
    )
# =============================================================
# CHECK CONSENT
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def onepageorder_check_consent(request, pk):

    if request.method not in (
        "GET",
        "POST",
    ):

        messages.warning(
            request,
            "Invalid request method.",
        )

        return redirect(
            "onepageorder_tracking_list"
        )

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
        ),
        pk=pk,
    )

    result = ConsentService.onepageorder_check_consent(
        session
    )

    if not result.get("success"):

        message = result.get(
            "message",
            "Unable to check consent.",
        )

        if isinstance(
            message,
            dict,
        ):

            message = (
                message.get("errorMessage")
                or message.get("message")
                or message.get("raw_response")
                or str(message)
            )

        messages.error(
            request,
            str(message),
        )

        return redirect(
            "onepageorder_tracking_setup",
            pk=session.order.pk,
        )

    # ---------------------------------------------------------
    # Service already updates TrackingSession.
    # Refresh DB and display the final state.
    # ---------------------------------------------------------

    session.refresh_from_db()

    status = session.status

    if (
        status == "waiting_location"
        and session.tracking_enabled
        and session.consent_received
    ):

        messages.success(
            request,
            (
                "Consent received. "
                "Live tracking has been activated."
            ),
        )

    elif status == "sms_sent":

        messages.warning(
            request,
            "Driver consent is still pending.",
        )

    elif status == "license_hold":

        messages.warning(
            request,
            "Tracking is currently on license hold.",
        )

    elif status == "expired":

        messages.warning(
            request,
            "Tracking consent has expired.",
        )

    elif status == "error":

        messages.error(
            request,
            (
                "Consent was processed, but "
                "tracking could not be activated."
            ),
        )

    else:

        messages.info(
            request,
            (
                "Consent status: "
                f"{session.get_status_display()}"
            ),
        )

    return redirect(
        "onepageorder_tracking_setup",
        pk=session.order.pk,
    )

# =============================================================
# DELETE TRACKING
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def delete_tracking(request, pk):

    if request.method != "POST":

        messages.warning(
            request,
            "Invalid request method.",
        )

        return redirect(
            "onepageorder_tracking_list"
        )

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
        ),
        pk=pk,
    )

    # ---------------------------------------------------------
    # Already deleted
    # ---------------------------------------------------------

    if session.status == "deleted":

        messages.info(
            request,
            "Tracking has already been deleted.",
        )

        return redirect(
            "onepageorder_tracking_list"
        )

    # ---------------------------------------------------------
    # Entity ID required
    # ---------------------------------------------------------

    if not session.entity_id:

        messages.warning(
            request,
            (
                "Tracking entity ID is missing. "
                "The Telenity tracking profile cannot be deleted."
            ),
        )

        return redirect(
            "onepageorder_tracking_setup",
            pk=session.order.pk,
        )

    # ---------------------------------------------------------
    # Delete from Telenity
    # ---------------------------------------------------------

    result = DeleteService.delete_tracking(
        session
    )

    if result.get("success"):

        # DeleteService already updates the session.
        session.refresh_from_db()

        messages.success(
            request,
            (
                "Tracking deleted successfully for "
                f"{session.order.trip_number}."
            ),
        )

        return redirect(
            "onepageorder_tracking_list"
        )

    message = result.get(
        "message",
        "Unable to delete tracking.",
    )

    if isinstance(
        message,
        dict,
    ):

        message = (
            message.get("errorMessage")
            or message.get("message")
            or message.get("raw_response")
            or str(message)
        )

    messages.error(
        request,
        str(message),
    )

    return redirect(
        "onepageorder_tracking_setup",
        pk=session.order.pk,
    )

# =============================================================
# TEST LOCATION API
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def test_location(request, pk):

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
        ),
        pk=pk,
    )

    driver_mobile = get_driver_mobile(
        session
    )

    if not driver_mobile:

        return JsonResponse(
            {
                "success": False,
                "message": (
                    "Driver mobile number "
                    "is not available."
                ),
            },
            status=400,
        )

    result = LocationService.get_location(
        driver_mobile
    )

    return JsonResponse(
        result,
        safe=False,
    )


# =============================================================
# LIVE VEHICLE PAGE
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def vehicle_live(request, pk):

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
            "order__customer",
            "order__tracking",
        ),
        pk=pk,
    )

    order = session.order

    latest_location = (
        session.locations
        .order_by("-received_at")
        .first()
    )

    return render(
        request,
        "onepageorders/vehicle_live.html",
        {
            "session": session,
            "order": order,
            "tracking": getattr(
                order,
                "tracking",
                None,
            ),
            "latest_location": latest_location,
            "locations": session.locations.all(),
        },
    )


# =============================================================
# LOCATION JSON
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def vehicle_live_location(request, pk):

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
        ),
        pk=pk,
    )

    order = session.order

    latest_location = (
        session.locations
        .order_by("-received_at")
        .first()
    )

    # ---------------------------------------------------------
    # No location yet
    # ---------------------------------------------------------

    if not latest_location:

        return JsonResponse(
            {
                "success": True,
                "tracking": True,
                "has_location": False,

                "trip_number": (
                    order.trip_number
                    or ""
                ),

                "vehicle_number": (
                    order.vehicle_number
                    or ""
                ),

                "driver_number": (
                    order.driver_number
                    or ""
                ),

                "tracking_reference": (
                    session.tracking_reference
                    or ""
                ),

                "status": (
                    session.get_status_display()
                ),

                "tracking_enabled": (
                    session.tracking_enabled
                ),

                "consent_received": (
                    session.consent_received
                ),

                "message": (
                    "Waiting for vehicle "
                    "location..."
                ),
            }
        )

    latitude = (
        float(latest_location.latitude)
        if latest_location.latitude is not None
        else None
    )

    longitude = (
        float(latest_location.longitude)
        if latest_location.longitude is not None
        else None
    )

    return JsonResponse(
        {
            "success": True,
            "tracking": True,
            "has_location": True,

            "trip_number": (
                order.trip_number
                or ""
            ),

            "vehicle_number": (
                order.vehicle_number
                or ""
            ),

            "driver_number": (
                order.driver_number
                or ""
            ),

            "tracking_reference": (
                session.tracking_reference
                or ""
            ),

            "status": (
                session.get_status_display()
            ),

            "tracking_enabled": (
                session.tracking_enabled
            ),

            "consent_received": (
                session.consent_received
            ),

            "latitude": latitude,
            "longitude": longitude,

            "accuracy": (
                latest_location.accuracy
            ),

            "location_name": (
                latest_location.location_name
                or ""
            ),

            "address": (
                latest_location.address
                or ""
            ),

            "location_status": (
                latest_location.location_status
                or ""
            ),

            "tracked": (
                latest_location.tracked
            ),

            "received_at": (
                latest_location.received_at.isoformat()
            ),

            "session_last_updated": (
                session.last_updated.isoformat()
                if session.last_updated
                else None
            ),
        }
    )


# =============================================================
# LOCATION HISTORY
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def vehicle_history(request, pk):

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
        ),
        pk=pk,
    )

    history = unique_location_history(
        session
    )

    return render(
        request,
        "onepageorders/history.html",
        {
            "session": session,
            "order": session.order,
            "history": history,
        },
    )


# =============================================================
# LIVE TRACKING HISTORY
# =============================================================

@login_required
def live_tracking_history(request, pk):

    return vehicle_history(
        request,
        pk,
    )


# =============================================================
# TRACKING HISTORY
# =============================================================
@login_required
def tracking_history(request, pk):
    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
            "order__customer",
        ),
        pk=pk,
    )
    history = (
        session.locations
        .all()
        .order_by("-received_at")
    )
    return render(
        request,
        "onepageorders/history.html",
        {
            "session": session,
            "order": session.order,
            "history": history,
            "history_count": history.count(),
        },
    )

# =============================================================
# REFRESH LOCATION
# pk = TRACKING SESSION PK
# =============================================================

@login_required
def refresh_location(request, pk):

    session = get_object_or_404(
        TrackingSession.objects.select_related(
            "order",
        ),
        pk=pk,
    )

    logger.info(
        "REFRESH LOCATION | "
        "session=%s order=%s mobile=%s "
        "entity=%s consent=%s status=%s",
        session.pk,
        session.order.trip_number,
        get_driver_mobile(session),
        session.entity_id,
        session.consent_received,
        session.status,
    )

    # ---------------------------------------------------------
    # 1. Consent
    # ---------------------------------------------------------

    if not session.consent_received:

        messages.warning(
            request,
            "Driver consent has not been approved yet.",
        )

        return redirect(
            "onepageordervehicle_live",
            pk=session.pk,
        )

    # ---------------------------------------------------------
    # 2. Start tracking if necessary
    # ---------------------------------------------------------

    if not session.tracking_enabled:

        modify_result = (
            ModifyService.start_tracking(
                session
            )
        )

        if not modify_result.get("success"):

            message = modify_result.get(
                "message",
                "Unable to start tracking.",
            )

            if isinstance(
                message,
                dict,
            ):
                message = str(message)

            messages.error(
                request,
                str(message),
            )

            return redirect(
                "onepageordervehicle_live",
                pk=session.pk,
            )

        session.refresh_from_db()

    # ---------------------------------------------------------
    # 3. Fetch latest location
    # ---------------------------------------------------------

    result = LocationService.fetch_location(
        session
    )

    logger.info(
        "REFRESH LOCATION RESULT | "
        "session=%s result=%s",
        session.pk,
        result,
    )

    # ---------------------------------------------------------
    # 4. User message
    # ---------------------------------------------------------

    if result.get("success"):

        if result.get("location_available"):

            messages.success(
                request,
                (
                    "Vehicle location updated "
                    "successfully."
                ),
            )

        else:

            messages.info(
                request,
                (
                    result.get(
                        "message",
                        (
                            "Tracking is active, "
                            "but the current location "
                            "is not available yet."
                        ),
                    )
                ),
            )

    else:

        message = result.get(
            "message",
            "Unable to retrieve vehicle location.",
        )

        if isinstance(
            message,
            dict,
        ):
            message = str(message)

        messages.error(
            request,
            str(message),
        )

    return redirect(
        "onepageordervehicle_live",
        pk=session.pk,
    )


# =============================================================
# TEST CONSENT AUTH
# =============================================================

@login_required
def test_consent_auth(request):

    result = (
        ConsentAuthService
        .get_consent_token()
    )

    return JsonResponse(
        result,
        safe=False,
    )


# =============================================================
# TEST TRACKING AUTH
# =============================================================

@login_required
def test_tracking_auth(request):

    result = (
        TrackingAuthService
        .get_tracking_token()
    )

    return JsonResponse(
        result,
        safe=False,
    )


# =============================================================
# API TOKEN STATUS
# =============================================================

@login_required
def api_token_status(request):
    tracking = (
        TrackingAuthService
        .get_tracking_token()
    )
    consent = (
        ConsentAuthService
        .get_consent_token()
    )

    return JsonResponse(
        {
            "tracking": tracking,
            "consent": consent,
        }
    )

from decimal import Decimal
from datetime import date
import os

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from .models import Order


# ============================================================
# PDF FONT
# ============================================================

FONT_DIR = os.path.join(
    settings.BASE_DIR,
    "static",
    "fonts",
)

DEJAVU_REGULAR = os.path.join(
    FONT_DIR,
    "DejaVuSans.ttf",
)

DEJAVU_BOLD = os.path.join(
    FONT_DIR,
    "DejaVuSans-Bold.ttf",
)

if os.path.exists(DEJAVU_REGULAR):

    try:
        pdfmetrics.registerFont(
            TTFont(
                "DejaVu",
                DEJAVU_REGULAR,
            )
        )
    except Exception:
        pass


if os.path.exists(DEJAVU_BOLD):

    try:
        pdfmetrics.registerFont(
            TTFont(
                "DejaVu-Bold",
                DEJAVU_BOLD,
            )
        )
    except Exception:
        pass


# ============================================================
# SAFE PDF FONT
# ============================================================

PDF_FONT = (
    "DejaVu"
    if os.path.exists(DEJAVU_REGULAR)
    else "Helvetica"
)

PDF_FONT_BOLD = (
    "DejaVu-Bold"
    if os.path.exists(DEJAVU_BOLD)
    else "Helvetica-Bold"
)


# ============================================================
# DECIMAL HELPER
# ============================================================

def safe_decimal(value):

    if value is None:
        return Decimal("0")

    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0")


def excel_datetime(value):

    if value is None:
        return None

    # Django DateTimeField
    if hasattr(value, "tzinfo"):

        if value.tzinfo is not None:

            value = timezone.localtime(
                value
            )

            value = value.replace(
                tzinfo=None
            )

    return value

# ============================================================
# REPORT DATA BUILDER
# ============================================================

def build_payment_report(request):

    query = request.GET.get(
        "q",
        ""
    ).strip()

    status_filter = request.GET.get(
        "status",
        ""
    ).strip().lower()

    from_date = request.GET.get(
        "from_date",
        ""
    ).strip()

    to_date = request.GET.get(
        "to_date",
        ""
    ).strip()

    # ========================================================
    # ORDERS
    # ========================================================

    orders = (
        Order.objects
        .select_related("customer")
        .order_by("-id")
    )

    # ========================================================
    # SEARCH
    # ========================================================

    if query:

        orders = orders.filter(

            Q(
                trip_number__icontains=query
            )

            |

            Q(
                vehicle_number__icontains=query
            )

            |

            Q(
                driver_number__icontains=query
            )

            |

            Q(
                origin__icontains=query
            )

            |

            Q(
                destination__icontains=query
            )

            |

            Q(
                customer__name__icontains=query
            )

        )

    # ========================================================
    # FIND AVAILABLE DATE FIELD
    # ========================================================

    existing_fields = {
        field.name
        for field in Order._meta.get_fields()
    }

    date_field = None

    for field_name in [
        "vehicle_place_date",
        "trip_date",
        "order_date",
        "created_at",
    ]:

        if field_name in existing_fields:

            date_field = field_name

            break

    # ========================================================
    # DATE FILTER
    # ========================================================

    if date_field:

        if from_date:

            orders = orders.filter(
                **{
                    f"{date_field}__date__gte": from_date
                }
            )

        if to_date:

            orders = orders.filter(
                **{
                    f"{date_field}__date__lte": to_date
                }
            )

    # ========================================================
    # REPORT
    # ========================================================

    report_rows = []

    total_billing = Decimal("0")
    total_paid = Decimal("0")
    total_advance = Decimal("0")
    total_balance = Decimal("0")
    total_recovery = Decimal("0")
    total_outstanding = Decimal("0")

    today = date.today()

    # ========================================================
    # LOOP
    # ========================================================

    for order in orders:

        # ----------------------------------------------------
        # CUSTOMER
        # ----------------------------------------------------

        customer_name = "-"

        if getattr(order, "customer", None):

            customer_name = (
                getattr(
                    order.customer,
                    "name",
                    None
                )
                or "-"
            )

        # ----------------------------------------------------
        # CONTACT
        # ----------------------------------------------------

        contact = (
            getattr(
                order,
                "customer_contact",
                None
            )
            or getattr(
                order.customer,
                "contact",
                None
            )
            if getattr(order, "customer", None)
            else "-"
        )

        contact = contact or "-"

        # ----------------------------------------------------
        # BILLING
        # ----------------------------------------------------

        billing_amount = (
            getattr(
                order,
                "total_rate",
                None
            )

            or getattr(
                order,
                "freight_amount",
                None
            )

            or getattr(
                order,
                "finalized_rate",
                None
            )

            or Decimal("0")
        )

        billing_amount = safe_decimal(
            billing_amount
        )

        # ----------------------------------------------------
        # ADVANCE
        # ----------------------------------------------------

        advance = safe_decimal(
            getattr(
                order,
                "advance",
                None
            )
        )

        # ----------------------------------------------------
        # BALANCE
        # ----------------------------------------------------

        balance_value = getattr(
            order,
            "balance",
            None
        )

        if balance_value is None:

            balance = (
                billing_amount
                - advance
            )

        else:

            balance = safe_decimal(
                balance_value
            )

        if balance < 0:

            balance = Decimal("0")

        # ----------------------------------------------------
        # PAID
        # ----------------------------------------------------

        paid_amount = (
            billing_amount
            - balance
        )

        if paid_amount < 0:

            paid_amount = Decimal("0")

        # ----------------------------------------------------
        # PAID TO PARTY
        # ----------------------------------------------------

        paid_to_party = (
            getattr(
                order,
                "paid_to_party",
                None
            )

            or getattr(
                order,
                "vehicle_paid",
                None
            )

            or Decimal("0")
        )

        paid_to_party = safe_decimal(
            paid_to_party
        )

        # ----------------------------------------------------
        # RECOVERY
        # ----------------------------------------------------

        recovery_amount = safe_decimal(
            getattr(
                order,
                "recovery_amount",
                None
            )
        )

        # ----------------------------------------------------
        # OUTSTANDING
        # ----------------------------------------------------

        outstanding_value = getattr(
            order,
            "outstanding",
            None
        )

        if outstanding_value is None:

            outstanding = (
                balance
                - recovery_amount
            )

        else:

            outstanding = safe_decimal(
                outstanding_value
            )

        if outstanding < 0:

            outstanding = Decimal("0")

        # ----------------------------------------------------
        # PROMISE DATE
        # ----------------------------------------------------

        promise_date = getattr(
            order,
            "promise_date",
            None
        )

        # ----------------------------------------------------
        # PAYMENT STATUS
        # ----------------------------------------------------

        if outstanding <= 0:

            payment_status = "paid"

        elif promise_date:

            try:

                promise_day = (
                    promise_date.date()
                    if hasattr(
                        promise_date,
                        "date"
                    )
                    else promise_date
                )

                if promise_day < today:

                    payment_status = "overdue"

                else:

                    payment_status = "promised"

            except Exception:

                payment_status = "due"

        else:

            payment_status = "due"

        # ----------------------------------------------------
        # STATUS FILTER
        # ----------------------------------------------------

        if status_filter:

            if status_filter != payment_status:

                continue

        # ----------------------------------------------------
        # TRIP DATE
        # ----------------------------------------------------

        trip_date = None

        if date_field:

            trip_date = getattr(
                order,
                date_field,
                None
            )

        # ----------------------------------------------------
        # TOTALS
        # ----------------------------------------------------

        total_billing += billing_amount
        total_paid += paid_amount
        total_advance += advance
        total_balance += balance
        total_recovery += recovery_amount
        total_outstanding += outstanding

        # ----------------------------------------------------
        # ROW
        # ----------------------------------------------------

        report_rows.append(
            {
                "id": order.pk,

                "trip_number": (
                    getattr(
                        order,
                        "trip_number",
                        None
                    )
                    or "-"
                ),

                "trip_date": trip_date,

                "customer": customer_name,

                "contact": contact,

                "billing_amount": billing_amount,

                "paid_to_party": paid_to_party,

                "paid_amount": paid_amount,

                "advance": advance,

                "balance": balance,

                "recovery_amount": recovery_amount,

                "outstanding": outstanding,

                "promise_date": promise_date,

                "payment_status": payment_status,

                "payment_terms": (
                    getattr(
                        order,
                        "payment_terms",
                        None
                    )
                    or "-"
                ),
            }
        )

    # ========================================================
    # COUNTS
    # ========================================================

    overdue_count = sum(
        1
        for row in report_rows
        if row["payment_status"] == "overdue"
    )

    due_count = sum(
        1
        for row in report_rows
        if row["payment_status"] == "due"
    )

    promised_count = sum(
        1
        for row in report_rows
        if row["payment_status"] == "promised"
    )

    paid_count = sum(
        1
        for row in report_rows
        if row["payment_status"] == "paid"
    )

    # ========================================================
    # RETURN
    # ========================================================

    return {
        "report_rows": report_rows,

        "total_billing": total_billing,
        "total_paid": total_paid,
        "total_advance": total_advance,
        "total_balance": total_balance,
        "total_recovery": total_recovery,
        "total_outstanding": total_outstanding,

        "overdue_count": overdue_count,
        "due_count": due_count,
        "promised_count": promised_count,
        "paid_count": paid_count,

        "query": query,
        "status_filter": status_filter,

        "from_date": from_date,
        "to_date": to_date,
    }


# ============================================================
# PAYMENT REPORT PAGE
# ============================================================

@login_required
def payment_report(request):

    context = build_payment_report(
        request
    )

    return render(
        request,
        "onepageorders/payment_report.html",
        context,
    )


# ============================================================
# PDF REPORT
# ============================================================

@login_required
def payment_report_pdf(request):

    context = build_payment_report(
        request
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="payment_report.pdf"'
    )

    # ========================================================
    # A4 LANDSCAPE
    # ========================================================

    doc = SimpleDocTemplate(

        response,

        pagesize=landscape(A4),

        leftMargin=8 * mm,
        rightMargin=8 * mm,

        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )

    # ========================================================
    # STYLES
    # ========================================================

    title_style = ParagraphStyle(

        "ReportTitle",

        fontName=PDF_FONT_BOLD,

        fontSize=17,

        leading=20,

        textColor=colors.HexColor(
            "#172033"
        ),

        alignment=TA_LEFT,

        spaceAfter=3,
    )

    subtitle_style = ParagraphStyle(

        "ReportSubtitle",

        fontName=PDF_FONT,

        fontSize=8,

        leading=10,

        textColor=colors.HexColor(
            "#667085"
        ),

        alignment=TA_LEFT,

        spaceAfter=10,
    )

    header_style = ParagraphStyle(

        "TableHeader",

        fontName=PDF_FONT_BOLD,

        fontSize=6.2,

        leading=7,

        textColor=colors.white,

        alignment=TA_CENTER,
    )

    text_style = ParagraphStyle(

        "TableText",

        fontName=PDF_FONT,

        fontSize=6.2,

        leading=7.5,

        textColor=colors.HexColor(
            "#344054"
        ),

        alignment=TA_LEFT,
    )

    number_style = ParagraphStyle(

        "TableNumber",

        fontName=PDF_FONT,

        fontSize=6.2,

        leading=7,

        textColor=colors.HexColor(
            "#344054"
        ),

        alignment=TA_RIGHT,
    )

    status_style = ParagraphStyle(

        "TableStatus",

        fontName=PDF_FONT_BOLD,

        fontSize=6.2,

        leading=7,

        alignment=TA_CENTER,
    )

    summary_label_style = ParagraphStyle(

        "SummaryLabel",

        fontName=PDF_FONT_BOLD,

        fontSize=7,

        leading=8,

        textColor=colors.HexColor(
            "#667085"
        ),

        alignment=TA_CENTER,
    )

    summary_value_style = ParagraphStyle(

        "SummaryValue",

        fontName=PDF_FONT_BOLD,

        fontSize=9,

        leading=11,

        textColor=colors.HexColor(
            "#172033"
        ),

        alignment=TA_CENTER,
    )

    elements = []

    # ========================================================
    # TITLE
    # ========================================================

    elements.append(
        Paragraph(
            "PAYMENT & RECOVERY REPORT",
            title_style,
        )
    )

    elements.append(
        Paragraph(
            "Overdue payments, customer recovery and outstanding balances",
            subtitle_style,
        )
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    summary_data = [

        [

            Paragraph(
                "TOTAL BILLING",
                summary_label_style,
            ),

            Paragraph(
                "TOTAL COLLECTED",
                summary_label_style,
            ),

            Paragraph(
                "BALANCE",
                summary_label_style,
            ),

            Paragraph(
                "RECOVERY",
                summary_label_style,
            ),

            Paragraph(
                "OUTSTANDING",
                summary_label_style,
            ),

        ],

        [

            Paragraph(
                f"₹ {context['total_billing']:,.2f}",
                summary_value_style,
            ),

            Paragraph(
                f"₹ {context['total_paid']:,.2f}",
                summary_value_style,
            ),

            Paragraph(
                f"₹ {context['total_balance']:,.2f}",
                summary_value_style,
            ),

            Paragraph(
                f"₹ {context['total_recovery']:,.2f}",
                summary_value_style,
            ),

            Paragraph(
                f"₹ {context['total_outstanding']:,.2f}",
                summary_value_style,
            ),

        ],
    ]

    summary_table = Table(

        summary_data,

        colWidths=[
            54 * mm,
            54 * mm,
            54 * mm,
            54 * mm,
            54 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor(
                    "#f8fafc"
                ),
            ),

            (
                "BOX",
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor(
                    "#e4e7ec"
                ),
            ),

            (
                "INNERGRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor(
                    "#e4e7ec"
                ),
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6,
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6,
            ),

        ])
    )

    elements.append(
        summary_table
    )

    elements.append(
        Spacer(
            1,
            7 * mm
        )
    )

    # ========================================================
    # TABLE HEADER
    # ========================================================

    table_data = [

        [

            Paragraph(
                "Trip No",
                header_style,
            ),

            Paragraph(
                "Date",
                header_style,
            ),

            Paragraph(
                "Customer",
                header_style,
            ),

            Paragraph(
                "Billing",
                header_style,
            ),

            Paragraph(
                "Paid",
                header_style,
            ),

            Paragraph(
                "Advance",
                header_style,
            ),

            Paragraph(
                "Balance",
                header_style,
            ),

            Paragraph(
                "Recovery",
                header_style,
            ),

            Paragraph(
                "Outstanding",
                header_style,
            ),

            Paragraph(
                "Promise<br/>Date",
                header_style,
            ),

            Paragraph(
                "Status",
                header_style,
            ),

        ]
    ]

    # ========================================================
    # DATA
    # ========================================================

    for row in context["report_rows"]:

        trip_date = excel_datetime(
            row["trip_date"]
        )

        promise_date = excel_datetime(
            row["promise_date"]
        )
        table_data.append([

            Paragraph(
                str(row["trip_number"]),
                text_style,
            ),

            Paragraph(
                (
                    trip_date.strftime(
                        "%d-%m-%Y"
                    )
                    if trip_date
                    else "-"
                ),
                text_style,
            ),

            # LONG CUSTOMER NAME WRAPS
            Paragraph(
                str(row["customer"]),
                text_style,
            ),

            Paragraph(
                f"₹ {row['billing_amount']:,.2f}",
                number_style,
            ),

            Paragraph(
                f"₹ {row['paid_amount']:,.2f}",
                number_style,
            ),

            Paragraph(
                f"₹ {row['advance']:,.2f}",
                number_style,
            ),

            Paragraph(
                f"₹ {row['balance']:,.2f}",
                number_style,
            ),

            Paragraph(
                f"₹ {row['recovery_amount']:,.2f}",
                number_style,
            ),

            Paragraph(
                f"₹ {row['outstanding']:,.2f}",
                number_style,
            ),

            Paragraph(
                (
                    promise_date.strftime(
                        "%d-%m-%Y"
                    )
                    if promise_date
                    else "-"
                ),
                text_style,
            ),

            Paragraph(
                row["payment_status"].title(),
                status_style,
            ),

        ])

    # ========================================================
    # TOTAL
    # ========================================================

    table_data.append([

        Paragraph(
            "TOTAL",
            ParagraphStyle(
                "Total",
                fontName=PDF_FONT_BOLD,
                fontSize=6.5,
                leading=7,
            ),
        ),

        "",
        "",

        Paragraph(
            f"₹ {context['total_billing']:,.2f}",
            number_style,
        ),

        Paragraph(
            f"₹ {context['total_paid']:,.2f}",
            number_style,
        ),

        Paragraph(
            f"₹ {context['total_advance']:,.2f}",
            number_style,
        ),

        Paragraph(
            f"₹ {context['total_balance']:,.2f}",
            number_style,
        ),

        Paragraph(
            f"₹ {context['total_recovery']:,.2f}",
            number_style,
        ),

        Paragraph(
            f"₹ {context['total_outstanding']:,.2f}",
            number_style,
        ),

        "",
        "",
    ])

    # ========================================================
    # COLUMN WIDTH
    #
    # TOTAL = 277 MM
    # ========================================================

    col_widths = [

        23 * mm,   # Trip

        21 * mm,   # Date

        39 * mm,   # Customer

        25 * mm,   # Billing

        25 * mm,   # Paid

        25 * mm,   # Advance

        25 * mm,   # Balance

        25 * mm,   # Recovery

        28 * mm,   # Outstanding

        22 * mm,   # Promise

        19 * mm,   # Status

    ]

    report_table = Table(

        table_data,

        colWidths=col_widths,

        repeatRows=1,

        hAlign="LEFT",
    )

    table_style = [

        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor(
                "#172033"
            ),
        ),

        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white,
        ),

        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.4,
            colors.HexColor(
                "#e4e7ec"
            ),
        ),

        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "MIDDLE",
        ),

        (
            "BACKGROUND",
            (0, -1),
            (-1, -1),
            colors.HexColor(
                "#f2f4f7"
            ),
        ),

        (
            "FONTNAME",
            (0, -1),
            (-1, -1),
            PDF_FONT_BOLD,
        ),

        (
            "LEFTPADDING",
            (0, 0),
            (-1, -1),
            4,
        ),

        (
            "RIGHTPADDING",
            (0, 0),
            (-1, -1),
            4,
        ),

        (
            "TOPPADDING",
            (0, 0),
            (-1, -1),
            4,
        ),

        (
            "BOTTOMPADDING",
            (0, 0),
            (-1, -1),
            4,
        ),
    ]

    # ========================================================
    # STATUS COLORS
    # ========================================================

    for index, row in enumerate(
        context["report_rows"],
        start=1,
    ):

        status = row[
            "payment_status"
        ]

        if status == "overdue":

            table_style.extend([

                (
                    "TEXTCOLOR",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#b42318"
                    ),
                ),

                (
                    "BACKGROUND",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#fef3f2"
                    ),
                ),

            ])

        elif status == "paid":

            table_style.extend([

                (
                    "TEXTCOLOR",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#087443"
                    ),
                ),

                (
                    "BACKGROUND",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#ecfdf3"
                    ),
                ),

            ])

        elif status == "promised":

            table_style.extend([

                (
                    "TEXTCOLOR",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#175cd3"
                    ),
                ),

                (
                    "BACKGROUND",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#eff8ff"
                    ),
                ),

            ])

        else:

            table_style.extend([

                (
                    "TEXTCOLOR",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#b54708"
                    ),
                ),

                (
                    "BACKGROUND",
                    (10, index),
                    (10, index),
                    colors.HexColor(
                        "#fff7ed"
                    ),
                ),

            ])

    report_table.setStyle(
        TableStyle(table_style)
    )

    elements.append(
        report_table
    )

    elements.append(
        Spacer(
            1,
            5 * mm
        )
    )

    # ========================================================
    # FOOTER
    # ========================================================

    footer_style = ParagraphStyle(

        "Footer",

        fontName=PDF_FONT,

        fontSize=6.5,

        textColor=colors.HexColor(
            "#98a2b3"
        ),

        alignment=TA_RIGHT,
    )

    elements.append(
        Paragraph(
            f"Generated on {date.today().strftime('%d-%m-%Y')}",
            footer_style,
        )
    )

    # ========================================================
    # BUILD
    # ========================================================

    doc.build(
        elements
    )

    return response


# ============================================================
# EXCEL REPORT
# ============================================================

from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)


# =========================================================
# EXCEL DATETIME HELPER
# =========================================================

def excel_safe_datetime(value):
    """
    Convert Django timezone-aware datetime to a
    timezone-naive datetime that Excel/openpyxl accepts.
    """

    if value is None:
        return None

    # DateTimeField
    if hasattr(value, "tzinfo"):

        if value.tzinfo is not None:

            # Convert to Django current timezone first
            value = timezone.localtime(value)

            # Remove timezone information
            value = value.replace(
                tzinfo=None
            )

    return value


# =========================================================
# PAYMENT REPORT - EXCEL
# =========================================================

@login_required
def payment_report_excel(request):

    context = build_payment_report(request)

    # =====================================================
    # WORKBOOK
    # =====================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Payment Report"


    # =====================================================
    # TITLE
    # =====================================================

    worksheet.merge_cells("A1:L1")

    worksheet["A1"] = (
        "PAYMENT & RECOVERY REPORT"
    )

    worksheet["A1"].font = Font(
        name="Calibri",
        size=16,
        bold=True,
        color="172033",
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 26


    # =====================================================
    # SUBTITLE
    # =====================================================

    worksheet.merge_cells("A2:L2")

    worksheet["A2"] = (
        "Overdue payments, customer recovery "
        "and outstanding balances"
    )

    worksheet["A2"].font = Font(
        name="Calibri",
        size=10,
        color="667085",
    )

    worksheet["A2"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )


    # =====================================================
    # SUMMARY
    # =====================================================

    summary_row = 4

    summary_headers = [

        "TOTAL BILLING",

        "TOTAL COLLECTED",

        "BALANCE",

        "RECOVERY",

        "OUTSTANDING",

    ]

    summary_values = [

        context["total_billing"],

        context["total_paid"],

        context["total_balance"],

        context["total_recovery"],

        context["total_outstanding"],

    ]

    summary_columns = [

        1,
        3,
        5,
        7,
        9,

    ]


    for column, header, value in zip(
        summary_columns,
        summary_headers,
        summary_values,
    ):

        # Header
        cell = worksheet.cell(
            row=summary_row,
            column=column,
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            size=9,
            color="667085",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="F8FAFC",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


        # Value
        value_cell = worksheet.cell(
            row=summary_row + 1,
            column=column,
        )

        value_cell.value = value

        value_cell.font = Font(
            bold=True,
            size=11,
            color="172033",
        )

        value_cell.number_format = (
            '₹ #,##0.00'
        )

        value_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


    # =====================================================
    # TABLE
    # =====================================================

    table_start = 7


    headers = [

        "Trip No",

        "Date",

        "Customer",

        "Contact",

        "Billing",

        "Paid",

        "Advance",

        "Balance",

        "Recovery",

        "Outstanding",

        "Promise Date",

        "Status",

    ]


    # =====================================================
    # TABLE HEADER
    # =====================================================

    for column, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=table_start,
            column=column,
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF",
            size=10,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="172033",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


    worksheet.row_dimensions[
        table_start
    ].height = 30


    # =====================================================
    # DATA
    # =====================================================

    current_row = table_start + 1


    for row in context["report_rows"]:

        # -------------------------------------------------
        # SAFE DATE / DATETIME
        # -------------------------------------------------

        trip_date = excel_safe_datetime(
            row.get("trip_date")
        )

        promise_date = excel_safe_datetime(
            row.get("promise_date")
        )


        # -------------------------------------------------
        # IMPORTANT:
        # 12 values for 12 columns
        # -------------------------------------------------

        values = [

            # 1
            row.get(
                "trip_number",
                "-"
            ),

            # 2
            trip_date,

            # 3
            row.get(
                "customer",
                "-"
            ),

            # 4
            row.get(
                "contact",
                "-"
            ),

            # 5
            row.get(
                "billing_amount",
                0
            ),

            # 6
            row.get(
                "paid_amount",
                0
            ),

            # 7
            row.get(
                "advance",
                0
            ),

            # 8
            row.get(
                "balance",
                0
            ),

            # 9
            row.get(
                "recovery_amount",
                0
            ),

            # 10
            row.get(
                "outstanding",
                0
            ),

            # 11
            promise_date,

            # 12
            row.get(
                "payment_status",
                "due",
            ).title(),

        ]


        # -------------------------------------------------
        # WRITE CELLS
        # -------------------------------------------------

        for column, value in enumerate(
            values,
            start=1,
        ):

            cell = worksheet.cell(
                row=current_row,
                column=column,
            )

            cell.value = value

            cell.font = Font(
                name="Calibri",
                size=9,
                color="344054",
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


            # -------------------------------------------------
            # DATE
            # -------------------------------------------------

            if column in [2, 11]:

                if value:

                    cell.number_format = (
                        "dd-mm-yyyy"
                    )

                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="top",
                    )


            # -------------------------------------------------
            # MONEY
            # -------------------------------------------------

            if column in [
                5,   # Billing
                6,   # Paid
                7,   # Advance
                8,   # Balance
                9,   # Recovery
                10,  # Outstanding
            ]:

                cell.number_format = (
                    '₹ #,##0.00'
                )

                cell.alignment = Alignment(
                    horizontal="right",
                    vertical="top",
                )


        # =================================================
        # STATUS
        # =================================================

        status = (
            row.get(
                "payment_status",
                "due",
            )
            .lower()
        )

        status_cell = worksheet.cell(
            row=current_row,
            column=12,
        )

        status_cell.value = (
            status.title()
        )

        status_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


        if status == "overdue":

            status_cell.font = Font(
                bold=True,
                color="B42318",
            )

            status_cell.fill = PatternFill(
                "solid",
                fgColor="FEF3F2",
            )


        elif status == "paid":

            status_cell.font = Font(
                bold=True,
                color="087443",
            )

            status_cell.fill = PatternFill(
                "solid",
                fgColor="ECFDF3",
            )


        elif status == "promised":

            status_cell.font = Font(
                bold=True,
                color="175CD3",
            )

            status_cell.fill = PatternFill(
                "solid",
                fgColor="EFF8FF",
            )


        else:

            status_cell.font = Font(
                bold=True,
                color="B54708",
            )

            status_cell.fill = PatternFill(
                "solid",
                fgColor="FFF7ED",
            )


        current_row += 1


    # =====================================================
    # TOTAL ROW
    # =====================================================

    total_row = current_row


    total_values = [

        # 1
        "TOTAL",

        # 2
        "",

        # 3
        "",

        # 4
        "",

        # 5
        context["total_billing"],

        # 6
        context["total_paid"],

        # 7
        context["total_advance"],

        # 8
        context["total_balance"],

        # 9
        context["total_recovery"],

        # 10
        context["total_outstanding"],

        # 11
        "",

        # 12
        "",

    ]


    for column, value in enumerate(
        total_values,
        start=1,
    ):

        cell = worksheet.cell(
            row=total_row,
            column=column,
        )

        cell.value = value

        cell.font = Font(
            bold=True,
            size=9,
            color="172033",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="F2F4F7",
        )

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )


        if column in [
            5,
            6,
            7,
            8,
            9,
            10,
        ]:

            cell.number_format = (
                '₹ #,##0.00'
            )

            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
            )


    # =====================================================
    # BORDERS
    # =====================================================

    thin_border = Border(

        left=Side(
            style="thin",
            color="D0D5DD",
        ),

        right=Side(
            style="thin",
            color="D0D5DD",
        ),

        top=Side(
            style="thin",
            color="D0D5DD",
        ),

        bottom=Side(
            style="thin",
            color="D0D5DD",
        ),

    )


    for row_cells in worksheet.iter_rows(

        min_row=table_start,

        max_row=total_row,

        min_col=1,

        max_col=12,

    ):

        for cell in row_cells:

            cell.border = thin_border


    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    widths = {

        "A": 16,    # Trip No

        "B": 14,    # Date

        "C": 32,    # Customer

        "D": 18,    # Contact

        "E": 17,    # Billing

        "F": 17,    # Paid

        "G": 17,    # Advance

        "H": 17,    # Balance

        "I": 17,    # Recovery

        "J": 19,    # Outstanding

        "K": 16,    # Promise Date

        "L": 14,    # Status

    }


    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width


    # =====================================================
    # ROW HEIGHT
    # =====================================================

    for row_number in range(
        table_start + 1,
        total_row + 1,
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 30


    # =====================================================
    # FREEZE
    # =====================================================

    worksheet.freeze_panes = "A8"


    # =====================================================
    # FILTER
    # =====================================================

    worksheet.auto_filter.ref = (
        f"A{table_start}:L{total_row}"
    )


    # =====================================================
    # PRINT SETTINGS
    # =====================================================

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.page_setup.orientation = (
        "landscape"
    )

    worksheet.page_setup.paperSize = (
        worksheet.PAPERSIZE_A4
    )

    worksheet.page_setup.fitToWidth = 1

    worksheet.page_setup.fitToHeight = 0

    worksheet.print_title_rows = (
        f"{table_start}:{table_start}"
    )

    worksheet.sheet_view.showGridLines = False


    # =====================================================
    # PAGE MARGINS
    # =====================================================

    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5


    # =====================================================
    # RESPONSE
    # =====================================================

    response = HttpResponse(

        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="payment_report.xlsx"'
    )

    workbook.save(
        response
    )
    return response